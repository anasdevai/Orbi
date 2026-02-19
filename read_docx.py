from docx import Document

# Read the Word document
doc = Document('Orbi_v4_Final.docx')

# Extract all text content
full_text = []
for para in doc.paragraphs:
    if para.text.strip():
        full_text.append(para.text)

# Also extract text from tables if any
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                full_text.append(cell.text)

# Write the content to a file with UTF-8 encoding
with open('document_content.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(full_text))

print("Content extracted successfully to document_content.txt")
